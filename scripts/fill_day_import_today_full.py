"""Build a rich, clearly labeled day-import workbook for 2026-09-05."""
from openpyxl import load_workbook
from pathlib import Path

SRC = Path(r"c:\Users\Smail\Downloads\oosm-day-import-template.xlsx")
if not SRC.exists():
    SRC = Path(r"c:\Users\Smail\Downloads\oosm-day-import-template (1).xlsx")
OUT = Path(r"c:\Users\Smail\Downloads\oosm-day-import-2026-09-05-FULL.xlsx")
OUT_SAMPLE = Path(r"c:\Users\Smail\Downloads\oosm-day-import-sample.xlsx")

DAY = "2026-09-05"
TANK = "Cuve-Demo-Today"
REGION = "Region-Demo-Sfax"
REGION2 = "Region-Demo-Kairouan"
PARCEL = "Parcel-Demo-Nord"
PARCEL2 = "Parcel-Demo-Sud"
STYPE_APP = "Apporteur"
STYPE_CLI = "Client"
CONT5 = "BIN-5L-DEMO"
CONT10 = "BIN-10L-DEMO"
SUP1 = "SUP-DEMO-01"
SUP2 = "SUP-DEMO-02"
SUP3 = "SUP-DEMO-03"


def clear_data(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def put(ws, r, values):
    for c, v in enumerate(values, start=1):
        ws.cell(r, c).value = v


wb = load_workbook(SRC)

# --- ImportMeta ---
ws = wb["ImportMeta"]
clear_data(ws)
put(ws, 2, ["businessDate", DAY])
put(ws, 3, ["timezone", "Africa/Tunis"])

# --- Masters ---
ws = wb["Regions"]
clear_data(ws)
put(ws, 2, [REGION, "Demo region Sfax — visible in reception forms"])
put(ws, 3, [REGION2, "Demo region Kairouan"])

ws = wb["Parcels"]
clear_data(ws)
put(ws, 2, [PARCEL, "Demo parcel North"])
put(ws, 3, [PARCEL2, "Demo parcel South"])

ws = wb["SupplierTypes"]
clear_data(ws)
put(ws, 2, [STYPE_APP, "Farmer / apporteur"])
put(ws, 3, [STYPE_CLI, "Buyer / client"])

ws = wb["QcRules"]
clear_data(ws)
put(ws, 2, ["Acidite", "Acidité", "true", "NUMERIC", 0, 3.3, "", "Oil acidity"])
put(ws, 3, ["Humidite", "Humidité", "false", "NUMERIC", 0, 70, "", "Olive humidity"])
put(ws, 4, ["MaturiteTxt", "Maturité", "false", "STRING", "", "", "Bonne", "Olive maturity (STRING)"])

ws = wb["Suppliers"]
clear_data(ws)
put(ws, 2, [SUP1, "Karim", "DemoBenAli", "97111001", "", REGION, STYPE_APP])
put(ws, 3, [SUP2, "Sonia", "DemoGharbi", "97111002", "", REGION, STYPE_APP])
put(ws, 4, [SUP3, "Nabil", "DemoClient", "97111003", "", REGION2, STYPE_CLI])

ws = wb["OilContainers"]
clear_data(ws)
put(ws, 2, [CONT5, CONT5, 5, 200, 2, 6])
put(ws, 3, [CONT10, CONT10, 10, 100, 3, 10])

ws = wb["StorageUnits"]
clear_data(ws)
put(ws, 2, [TANK, TANK])
put(ws, 3, ["Cuve-Demo-Reserve", "Cuve-Demo-Reserve"])

# --- Receptions (operation types match UI lists) ---
# cols: externalRef, deliveryType, oliveOilType, varietyName, operationType, supplierKey, regionName, parcelName,
#       poidsNet, oilQuantity, unitPrice, storageUnitKey, description
ws = wb["Receptions"]
clear_data(ws)
put(ws, 2, [
    "R-MILL-DEMO-01", "OLIVE", "OC", "Chemlali", "SIMPLE_RECEPTION", SUP1, REGION, PARCEL,
    3200, 480, "", "", "DEMO TODAY — Private milling (SIMPLE_RECEPTION) — look under Trituration / Completed"
])
put(ws, 3, [
    "R-MILL-DEMO-02", "OLIVE", "OB", "Chetoui", "SIMPLE_RECEPTION", SUP2, REGION, PARCEL2,
    1800, 250, "", "", "DEMO TODAY — 2nd private milling lot"
])
put(ws, 4, [
    "R-OLVBUY-DEMO-01", "OLIVE", "OC", "Oueslati", "OLIVE_PURCHASE", SUP1, REGION, PARCEL,
    1500, 200, 2.5, "", "DEMO TODAY — Olive purchase — list OLIVE_PURCHASE / Completed"
])
put(ws, 5, [
    "R-OILBUY-DEMO-01", "OIL", "OC", "Chemlali oil", "OIL_PURCHASE", SUP3, REGION2, "",
    0, 220, 11.5, TANK, "DEMO TODAY — Oil purchase INTO Cuve-Demo-Today — list OIL_PURCHASE / IN_STOCK"
])
put(ws, 6, [
    "R-OILBUY-DEMO-02", "OIL", "OC", "Chetoui oil", "OIL_PURCHASE", SUP3, REGION2, "",
    0, 80, 12.0, TANK, "DEMO TODAY — 2nd oil purchase into same tank"
])
put(ws, 7, [
    "R-EXCH-DEMO-01", "OLIVE", "OB", "Zalmati", "EXCHANGE", SUP2, REGION, PARCEL,
    900, 120, "", "", "DEMO TODAY — Exchange operation"
])
put(ws, 8, [
    "R-BASE-DEMO-01", "OLIVE", "OC", "Chemlali", "BASE", SUP1, REGION, PARCEL2,
    2100, 300, "", "", "DEMO TODAY — BASE (cooperative) — list shows PROD_READY"
])

# --- QC ---
ws = wb["QcResults"]
clear_data(ws)
put(ws, 2, ["R-MILL-DEMO-01", "Humidite", "48", "false"])
put(ws, 3, ["R-MILL-DEMO-01", "MaturiteTxt", "Bonne", "false"])
put(ws, 4, ["R-MILL-DEMO-02", "Humidite", "52", "false"])
put(ws, 5, ["R-OLVBUY-DEMO-01", "Humidite", "45", "false"])
put(ws, 6, ["R-OILBUY-DEMO-01", "Acidite", "0.32", "true"])
put(ws, 7, ["R-OILBUY-DEMO-02", "Acidite", "0.41", "true"])
put(ws, 8, ["R-EXCH-DEMO-01", "Humidite", "50", "false"])
put(ws, 9, ["R-BASE-DEMO-01", "Humidite", "47", "false"])

# --- Payments (priced rows) ---
ws = wb["Payments"]
clear_data(ws)
put(ws, 2, ["R-OLVBUY-DEMO-01", 800, "CASH"])
put(ws, 3, ["R-OILBUY-DEMO-01", 1500, "CASH"])
put(ws, 4, ["R-OILBUY-DEMO-02", 500, "CHECK"])

# --- Oil sales (need tank volume >= qty at dry-run time) ---
# After clean DB create tank with >= 50 L OR commit oil buys first in a prior file.
# Same-file: dry-run checks CURRENT tank volume only — keep sale qty modest.
ws = wb["OilSales"]
clear_data(ws)
put(ws, 2, [
    "S-DEMO-01", "INV-DEMO-20260905-01", SUP3, TANK, 15, 16.5, "TND", "CASH",
    "EXTRA_VIRGIN", 247.5, "DEMO TODAY — Oil sale 15L — Finance → Oil sales"
])
put(ws, 3, [
    "S-DEMO-02", "INV-DEMO-20260905-02", SUP3, TANK, 10, 17.0, "TND", "CASH",
    "VIRGIN", 170, "DEMO TODAY — Oil sale 10L + containers"
])

ws = wb["OilSaleContainers"]
clear_data(ws)
put(ws, 2, ["S-DEMO-02", CONT5, 2])
put(ws, 3, ["S-DEMO-02", CONT10, 1])

# --- Expenses ---
ws = wb["Expenses"]
clear_data(ws)
put(ws, 2, ["EXP-DEMO-01", 120.5, "Fuel diesel", "Diesel", "OTHER", "CASH", "Station Demo", "FUEL-DEMO-1", "DEMO TODAY expense fuel"])
put(ws, 3, ["EXP-DEMO-02", 85.0, "Spare parts", "Maintenance", "OTHER", "CASH", "Atelier Demo", "PARTS-DEMO-1", "DEMO TODAY expense parts"])
put(ws, 4, ["EXP-DEMO-03", 40.0, "Office supplies", "Admin", "OTHER", "CASH", "Librairie", "ADM-DEMO-1", "DEMO TODAY expense admin"])

# --- README ---
ws = wb["README"]
ws["A1"] = f"FULL DEMO day {DAY} — search descriptions containing DEMO TODAY after commit."
ws["A2"] = "BEFORE import: create storage unit named exactly Cuve-Demo-Today with currentVolume >= 25 (for oil sales dry-run)."
ws["A3"] = "Where to look after commit (restart BE if finalize statuses were just added):"
ws["A4"] = "  • Reception → Olive / SIMPLE_RECEPTION (Completed): R-MILL-DEMO-01/02"
ws["A5"] = "  • Reception → Olive purchase (Completed): R-OLVBUY-DEMO-01"
ws["A6"] = "  • Reception → Oil purchase (IN_STOCK): R-OILBUY-DEMO-01/02"
ws["A7"] = "  • Exchange / BASE lists: R-EXCH-DEMO-01, R-BASE-DEMO-01"
ws["A8"] = "  • Finance → Oil sales: INV-DEMO-20260905-01/02 | Expenses: EXP-DEMO-*"
ws["A9"] = "  • Suppliers: Karim/Sonia/Nabil Demo* | Tank: Cuve-Demo-Today | Containers: BIN-*-DEMO"
ws["A10"] = "Dry-run first. OIL rows must use operationType OIL_PURCHASE (not OLIVE_PURCHASE)."

wb.save(OUT)
try:
    wb.save(OUT_SAMPLE)
except PermissionError:
    print("sample locked — skipped overwrite")
print("Wrote", OUT)
print("Tank required:", TANK, ">= 25 L currentVolume before dry-run (sales)")
print("Receptions: 7 | Sales: 2 | Expenses: 3 | Suppliers: 3")
