from openpyxl import load_workbook

path = r"c:\Users\Smail\Downloads\oosm-day-import-sample.xlsx"
wb = load_workbook(path)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} ({ws.max_row}x{ws.max_column}) ===")
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None and str(v).strip() != "" for v in vals):
            print(r, vals)
