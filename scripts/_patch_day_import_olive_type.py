"""Insert oliveOilType into Receptions and fill OC/OB for the FULL demo workbook."""
from openpyxl import load_workbook
from pathlib import Path

SRC = Path(r"c:\Users\Smail\Downloads\oosm-day-import-2026-09-05-FULL.xlsx")
OUT = SRC

wb = load_workbook(SRC)
ws = wb["Receptions"]
headers = [c.value for c in ws[1]]
print("before:", headers)

if "oliveOilType" not in headers:
    # Insert after deliveryType (col 2) -> new col 3
    ws.insert_cols(3)
    ws.cell(1, 3).value = "oliveOilType"
    headers = [c.value for c in ws[1]]

type_col = headers.index("oliveOilType") + 1
delivery_col = headers.index("deliveryType") + 1

# Prefer OC for demo lots (matches UI screenshot style); alternate OB on 2nd olive row.
olive_i = 0
for r in range(2, ws.max_row + 1):
    delivery = (ws.cell(r, delivery_col).value or "").strip().upper()
    if not delivery:
        continue
    if delivery == "OLIVE":
        ws.cell(r, type_col).value = "OC" if olive_i % 2 == 0 else "OB"
        olive_i += 1
    else:
        ws.cell(r, type_col).value = "OC"  # oil uses same enum codes

print("after:", [c.value for c in ws[1]])
for r in range(2, min(ws.max_row, 9) + 1):
    print("row", r, [ws.cell(r, c).value for c in range(1, 6)])

wb.save(OUT)
print("Saved", OUT)
