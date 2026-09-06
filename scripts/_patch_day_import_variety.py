"""Add varietyName to Receptions (generic OLIVE_VARIETY / OIL_VARIETY)."""
from openpyxl import load_workbook
from pathlib import Path

SRC = Path(r"c:\Users\Smail\Downloads\oosm-day-import-2026-09-05-FULL.xlsx")
wb = load_workbook(SRC)
ws = wb["Receptions"]
headers = [c.value for c in ws[1]]
print("before", headers)

if "varietyName" not in headers:
    # Insert after oliveOilType if present, else after deliveryType
    if "oliveOilType" in headers:
        insert_at = headers.index("oliveOilType") + 2  # 1-based excel col after oliveOilType
    else:
        insert_at = headers.index("deliveryType") + 2
    ws.insert_cols(insert_at)
    ws.cell(1, insert_at).value = "varietyName"
    headers = [c.value for c in ws[1]]

variety_col = headers.index("varietyName") + 1
delivery_col = headers.index("deliveryType") + 1

olive_i = 0
olive_vars = ["Chemlali", "Chetoui", "Oueslati", "Zalmati"]
oil_vars = ["Chemlali oil", "Chetoui oil"]
oil_i = 0
for r in range(2, ws.max_row + 1):
    delivery = (ws.cell(r, delivery_col).value or "").strip().upper()
    if not delivery:
        continue
    if delivery == "OLIVE":
        ws.cell(r, variety_col).value = olive_vars[olive_i % len(olive_vars)]
        olive_i += 1
    else:
        ws.cell(r, variety_col).value = oil_vars[oil_i % len(oil_vars)]
        oil_i += 1

print("after", [c.value for c in ws[1]])
for r in range(2, min(ws.max_row, 9) + 1):
    print("row", r, [ws.cell(r, c).value for c in range(1, 7)])
wb.save(SRC)
print("Saved", SRC)
