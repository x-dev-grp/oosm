"""Fill oosm-day-import-sample.xlsx with a coherent test day and unique refs."""
from openpyxl import load_workbook
from copy import copy
from pathlib import Path

SRC = Path(r"c:\Users\Smail\Downloads\oosm-day-import-sample.xlsx")
OUT = Path(r"c:\Users\Smail\Downloads\oosm-day-import-manual.xlsx")
# also overwrite the user's sample path for convenience
OUT_SAMPLE = Path(r"c:\Users\Smail\Downloads\oosm-day-import-sample.xlsx")

DAY = "2026-09-05"
TANK = "Cuve-1"
REGION = "TestRegion-Sfax"
PARCEL = "TestParcel-A"
STYPE = "Apporteur"
SUP = "SUP-TEST-01"
CONT = "BIN-5L"
R_OIL = "R-OIL-20260905-02"
R_OLIVE = "R-OLV-20260905-02"
SALE = "S-20260905-02"
INV = "INV-TEST-20260905-02"
EXP = "EXP-20260905-02"
RULE = "Acidite"


def set_row(ws, row_idx, values):
    for col, val in enumerate(values, start=1):
        ws.cell(row_idx, col).value = val


def clear_data_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


wb = load_workbook(SRC)

# ImportMeta
meta = wb["ImportMeta"]
clear_data_rows(meta)
set_row(meta, 2, ["businessDate", DAY])
set_row(meta, 3, ["timezone", "Africa/Tunis"])

# Masters
ws = wb["Regions"]
clear_data_rows(ws)
set_row(ws, 2, [REGION, "Sfax test region for day import"])
set_row(ws, 3, ["TestRegion-SidiBouzid", "Second region"])

ws = wb["Parcels"]
clear_data_rows(ws)
set_row(ws, 2, [PARCEL, "Parcel A"])
set_row(ws, 3, ["TestParcel-B", "Parcel B"])

ws = wb["SupplierTypes"]
clear_data_rows(ws)
set_row(ws, 2, [STYPE, "Default apporteur"])
set_row(ws, 3, ["Client", "Buyer / client"])

ws = wb["QcRules"]
clear_data_rows(ws)
set_row(ws, 2, [RULE, "Acidité", "true", "NUMERIC", 0, 3.3, "", "Oil acidity"])
set_row(ws, 3, ["Humidite", "Humidité", "false", "NUMERIC", 0, 70, "", "Olive humidity"])

ws = wb["Suppliers"]
clear_data_rows(ws)
set_row(ws, 2, [SUP, "Mohamed", "Trabelsi", "98111222", "", REGION, STYPE])
set_row(ws, 3, ["SUP-TEST-02", "Fatma", "Gharbi", "98333444", "", REGION, "Client"])

ws = wb["OilContainers"]
clear_data_rows(ws)
set_row(ws, 2, [CONT, CONT, 5, 100, 2, 5])
set_row(ws, 3, ["BIN-10L", "BIN-10L", 10, 50, 3, 8])

ws = wb["StorageUnits"]
clear_data_rows(ws)
set_row(ws, 2, [TANK, TANK])

ws = wb["Receptions"]
clear_data_rows(ws)
# OIL purchase with stock into Cuve-1
set_row(ws, 2, [R_OIL, "OIL", "OLIVE_PURCHASE", SUP, REGION, PARCEL, 1200, 180, 11.5, TANK, "Test oil purchase day import"])
# OLIVE private milling (no tank required)
set_row(ws, 3, [R_OLIVE, "OLIVE", "SIMPLE_RECEPTION", SUP, REGION, PARCEL, 2500, "", "", "", "Test olive private milling"])

ws = wb["QcResults"]
clear_data_rows(ws)
set_row(ws, 2, [R_OIL, RULE, "0.35", "true"])
set_row(ws, 3, [R_OLIVE, "Humidite", "42", "false"])

ws = wb["Payments"]
clear_data_rows(ws)
set_row(ws, 2, [R_OIL, 500, "CASH"])

ws = wb["OilSales"]
clear_data_rows(ws)
set_row(ws, 2, [SALE, INV, SUP, TANK, 25, 16, "TND", "CASH", "EXTRA_VIRGIN", 400, "Test sale after seed volume"])

ws = wb["OilSaleContainers"]
clear_data_rows(ws)
set_row(ws, 2, [SALE, CONT, 2])

ws = wb["Expenses"]
clear_data_rows(ws)
set_row(ws, 2, [EXP, 75.5, "Fuel", "Diesel", "OTHER", "CASH", "Total Energies", "FUEL-0905", "Day import test expense"])

readme = wb["README"]
readme["A1"] = f"Test day file for {DAY} — unique refs R-OIL/R-OLV/S/EXP-20260905. Tank {TANK} must exist in DB."
readme["A2"] = "Dry-run first; commit only with zero ERROR. Masters create-if-missing; duplicates skip."
readme["A3"] = f"Refs: {R_OIL}, {R_OLIVE}, {SALE}/{INV}, {EXP}. Container {CONT}."

wb.save(OUT)
wb.save(OUT_SAMPLE)
print(f"Wrote {OUT}")
print(f"Updated {OUT_SAMPLE}")
print("Refs:", R_OIL, R_OLIVE, SALE, INV, EXP, "tank=", TANK)
